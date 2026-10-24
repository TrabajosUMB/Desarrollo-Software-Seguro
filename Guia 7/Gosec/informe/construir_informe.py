# -*- coding: utf-8 -*-
"""Genera el informe Word (estilo UMB, APA 7) de la Guía 7 sobre gosec.

Las cifras salen de los JSON reales de ../resultados y se comprueban con `assert`; si no coinciden, no se genera el documento.
Requiere el generador institucional (estilos reales de TallerBI.docx) y el módulo de correcciones del skill umb-documento:
ambos viven fuera de este repositorio, en el directorio del usuario (ver GEN y REFS).

Uso:  py -I informe/construir_informe.py        (desde la carpeta Gosec)
      py -I <skill umb-documento>/references/exportar_pdf.py informe/Informe_Gosec_Guia7.docx
"""
import collections
import json
import pathlib
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx.shared import Inches, Pt
from matplotlib.patches import Patch
from PIL import Image

HOME = pathlib.Path.home()
GEN = HOME / "Documents" / "Universidad" / "Katherin" / "2026-02" / "Sistemas Transaccionales" / "herramientas_doc" / "build"
REFS = HOME / ".claude" / "skills" / "umb-documento" / "references"
sys.path.insert(0, str(GEN))
sys.path.insert(0, str(REFS))
from constructor import Documento  # noqa: E402
from corregir_docx import agregar_referencia, cabecera_azul, crear_estilo_codigo  # noqa: E402

RAIZ = pathlib.Path(__file__).resolve().parent.parent
RES = RAIZ / "resultados"
IMG = pathlib.Path(__file__).resolve().parent / "img"
SALIDA = pathlib.Path(__file__).resolve().parent / "Informe_Gosec_Guia7.docx"

TITULO = "Análisis estático de seguridad (SAST) con gosec"
SUBTITULO = "Informe de la Guía 7: instalación, hallazgos y experiencia de uso"
AUTORES = ("Cobos, Dilan Steven", "Rodríguez Ángel, Santiago", "Alarcón Ospina, Víctor Alfonso")

# ---------------------------------------------------------------- datos reales
govwa = json.loads((RES / "govwa.json").read_text(encoding="utf-8"))
vuln = json.loads((RES / "ejemplo_vulnerable.json").read_text(encoding="utf-8"))
corr = json.loads((RES / "ejemplo_corregido.json").read_text(encoding="utf-8"))
G = govwa["Issues"]
SEV = collections.Counter(i["severity"] for i in G)
CONF = collections.Counter(i["confidence"] for i in G)
MATRIZ = collections.Counter((i["severity"], i["confidence"]) for i in G)
POR_REGLA = collections.Counter(i["rule_id"] for i in G)
V_ACTIVOS = [i for i in vuln["Issues"] if not i.get("suppressions")]
C_ACTIVOS = [i for i in corr["Issues"] if not i.get("suppressions")]
C_SUPRIM = [i for i in corr["Issues"] if i.get("suppressions")]

assert govwa["Stats"] == {"files": 20, "lines": 1569, "nosec": 0, "found": 35}, govwa["Stats"]
assert (SEV["HIGH"], SEV["MEDIUM"], SEV["LOW"]) == (1, 16, 18)
assert (CONF["HIGH"], CONF["LOW"]) == (28, 7)
assert MATRIZ == {("HIGH", "HIGH"): 1, ("MEDIUM", "HIGH"): 9, ("MEDIUM", "LOW"): 7, ("LOW", "HIGH"): 18}
assert len(POR_REGLA) == 9
assert len(V_ACTIVOS) == 17 and not C_ACTIVOS and len(C_SUPRIM) == 2
assert collections.Counter(i["severity"] for i in V_ACTIVOS) == {"HIGH": 6, "MEDIUM": 9, "LOW": 2}
assert len({i["rule_id"] for i in V_ACTIVOS}) == 15

NOMBRE_REGLA = {
    "G701": "Inyección SQL", "G203": "HTML sin escapar", "G401": "Criptografía débil (MD5)", "G501": "Importa crypto/md5",
    "G124": "Cookie sin atributos seguros", "G201": "SQL armado con formato", "G112": "Sin ReadHeaderTimeout",
    "G706": "Inyección en registros", "G104": "Error sin manejar",
}
ORDEN_SEV = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
COLOR_SEV = {"HIGH": "#1F497D", "MEDIUM": "#6E94C8", "LOW": "#BFD1E8"}
TEXTO_SEV = {"HIGH": "Alta", "MEDIUM": "Media", "LOW": "Baja"}
REGLA_SEV = {i["rule_id"]: (i["severity"], i["confidence"], i["cwe"]["id"]) for i in G}


def figura_hallazgos(ruta):
    """Barras horizontales por regla, coloreadas por severidad, a partir de los hallazgos reales."""
    font_manager.fontManager.addfont(r"C:\Windows\Fonts\times.ttf")
    plt.rcParams["font.family"] = "Times New Roman"
    reglas = sorted(POR_REGLA, key=lambda r: (ORDEN_SEV[REGLA_SEV[r][0]], -POR_REGLA[r], r))
    etiquetas = [f"{r} · {NOMBRE_REGLA[r]}" for r in reglas]
    valores = [POR_REGLA[r] for r in reglas]
    colores = [COLOR_SEV[REGLA_SEV[r][0]] for r in reglas]
    fig, ax = plt.subplots(figsize=(6.5, 3.5), dpi=250)
    ax.barh(range(len(reglas)), valores, color=colores, edgecolor="#1F497D", linewidth=0.6, height=0.68)
    ax.set_yticks(range(len(reglas)), etiquetas, fontsize=9)
    ax.invert_yaxis()
    for y, v in enumerate(valores):
        ax.text(v + 0.2, y, str(v), va="center", fontsize=9, fontweight="bold")
    ax.set_xlabel("Cantidad de hallazgos", fontsize=9)
    ax.set_xlim(0, max(valores) + 1.5)
    ax.tick_params(axis="x", labelsize=8)
    ax.grid(axis="x", color="#D9D9D9", linewidth=0.5)
    ax.set_axisbelow(True)
    for lado in ("top", "right"):
        ax.spines[lado].set_visible(False)
    ax.legend(handles=[Patch(facecolor=COLOR_SEV[s], edgecolor="#1F497D", label=f"Severidad {TEXTO_SEV[s].lower()}") for s in ORDEN_SEV],
              fontsize=8.5, frameon=False, loc="upper center", bbox_to_anchor=(0.4, -0.2), ncol=3)
    fig.tight_layout()
    fig.savefig(ruta, facecolor="white")
    plt.close(fig)


def figura_informe_html(origen, ruta):
    """Recorte de la captura real del informe HTML: filtros, resumen y los dos primeros hallazgos (sin barra de desplazamiento)."""
    with Image.open(origen) as im:
        im.crop((120, 80, 2080, 1240)).save(ruta)


IMG.mkdir(exist_ok=True)
figura_hallazgos(IMG / "figura2_hallazgos_por_regla.png")
figura_informe_html(IMG / "figura1_informe_html.png", IMG / "figura1_recorte.png")


# ---------------------------------------------------------------- utilidades de documento
def tab(doc, *args, **kw):
    t = doc.tabla(*args, **kw)
    cabecera_azul(t)
    return t


def figura_ancho(doc, clave, titulo, archivo, nota, ancho_in):
    """Como Documento.figura, pero con un ancho menor al útil (centrada); sirve para que el bloque rótulo + imagen + nota
    quepa en el resto de una página en lugar de saltar a la siguiente y dejar un hueco."""
    etiqueta = doc.etiquetas.get(f"F:{clave}") or doc.registrar_figura(clave)
    doc._rotulo("TITULO FIGURA", etiqueta, titulo)
    p = doc.d.add_paragraph(style="Portada")
    pf = p.paragraph_format
    pf.line_spacing, pf.space_after, pf.keep_with_next = 1.0, Pt(4), True
    p.add_run().add_picture(str(archivo), width=Inches(ancho_in))
    doc._nota(nota)


def codigo(doc, lineas):
    """Bloque de código en el estilo `Codigo` (una línea por párrafo; el bloque no se parte entre páginas)."""
    for k, linea in enumerate(lineas):
        p = doc.d.add_paragraph(style="Codigo")
        p.add_run(linea)
        p.paragraph_format.keep_with_next = k < len(lineas) - 1
        p.paragraph_format.keep_together = True
        if k == len(lineas) - 1:
            p.paragraph_format.space_after = Pt(10)


# ---------------------------------------------------------------- portada y listas
def portada(doc):
    doc.fijar_propiedades(f"{TITULO}: {SUBTITULO}", AUTORES)
    doc.portada_logo(3.0)
    doc.portada_linea()
    doc.portada_linea(TITULO, negrita=True)
    doc.portada_linea(SUBTITULO)
    doc.portada_linea()
    for linea in (*AUTORES, "Universidad Manuela Beltrán", "Ingeniería de Software", "Facultad de Ingeniería",
                  "Desarrollo de Software Seguro", "Bogotá, Colombia", "2026"):
        doc.portada_linea(linea)
    doc.salto_pagina()


def listas(doc):
    doc.rotulo_centrado("Tabla de contenido")
    doc.campo_toc('TOC \\o "1-3" \\h \\z \\u')
    doc.rotulo_centrado("Lista de tablas", salto_antes=True)
    doc.campo_toc('TOC \\h \\z \\t "TITULO TABLA,1"')
    doc.rotulo_centrado("Lista de figuras", salto_antes=True)
    doc.campo_toc('TOC \\h \\z \\t "TITULO FIGURA,1"')


# ---------------------------------------------------------------- contenido
def introduccion(doc):
    doc.titulo(1, "Introducción", salto_antes=True)
    doc.parrafo(
        "Las pruebas estáticas de seguridad de aplicaciones (SAST, por las siglas en inglés de Static Application Security Testing) "
        "revisan el código fuente sin ejecutar la aplicación y buscan patrones asociados a vulnerabilidades conocidas (OWASP Foundation, "
        "s. f.). Su principal ventaja es que permiten detectar problemas desde que el código se escribe, cuando corregirlos es más sencillo.")
    doc.parrafo(
        "Este informe documenta la Guía 7 de la asignatura Desarrollo de Software Seguro, cuyo tema asignado al grupo es gosec, un "
        "analizador SAST para el lenguaje Go mantenido por el proyecto securego (securego, s. f.). El objetivo fue instalar y configurar "
        "la herramienta, analizar una aplicación vulnerable y código propio, documentar las vulnerabilidades halladas con su severidad y "
        "sus medidas de mitigación, y valorar la experiencia de uso en cuanto a instalación, facilidad, resultados e interfaz.")
    doc.parrafo(
        "El documento se organiza en ocho secciones. La sección 1 presenta el marco conceptual; la 2, la configuración del entorno; la 3, "
        "la metodología; la 4, los resultados sobre la aplicación vulnerable; la 5, el ciclo de detección y corrección con código propio; "
        "la 6, los límites de la herramienta; la 7, la experiencia de uso, y la 8, las recomendaciones. Todas las cifras provienen de "
        "ejecuciones reales de gosec conservadas en la carpeta de resultados del trabajo, por lo que pueden repetirse con los scripts que "
        "se describen en el Anexo A.")


def seccion1(doc):
    doc.titulo(1, "1. Marco conceptual")
    doc.titulo(2, "1.1 Análisis estático de seguridad")
    doc.parrafo(
        "El análisis estático examina el código sin ejecutarlo, de forma comparable a revisar los planos de una casa antes de construirla. "
        "Su alcance es el código que recibe: no observa el comportamiento de la aplicación en funcionamiento, ni los componentes de "
        "terceros, ni la configuración del servidor. Por esa razón se complementa con las pruebas dinámicas (DAST), que atacan la "
        "aplicación mientras se ejecuta, y con el análisis de composición de software (SCA), que revisa las dependencias.")
    doc.titulo(2, "1.2 ¿Qué es gosec?")
    doc.parrafo(
        "gosec es un analizador de seguridad de código abierto para el lenguaje Go, publicado bajo licencia Apache 2.0 (securego, s. f.). "
        "Examina el árbol sintáctico (AST) y la representación intermedia SSA del programa y aplica un conjunto de reglas, cada una con "
        "un identificador propio (de G101 a G710) y asociada a una debilidad del catálogo Common Weakness Enumeration, CWE (MITRE, s. f.). "
        "Las reglas de la serie G70x añaden análisis de flujo de datos, que sigue un dato desde la entrada del usuario hasta una función "
        "peligrosa. La {T:gosec} resume las características de la versión utilizada.")
    tab(doc, "gosec", "Características de gosec v2.30.0", ["Característica", "Descripción"], [
        ("Lenguaje analizado", "Go (solo archivos .go)"),
        ("Licencia", "Apache 2.0"),
        ("Reglas", "62, de G101 a G710, cada una con su CWE"),
        ("Técnicas", "Árbol sintáctico (AST), representación SSA y análisis de flujo de datos (G701 a G710)"),
        ("Filtros", "Por severidad y confianza, e inclusión o exclusión de reglas"),
        ("Formatos de salida", "text, json, yaml, csv, junit-xml, html, sonarqube, golint y sarif"),
        ("Integración continua", "GitHub Action oficial (securego/gosec)"),
        ("Requisito", "Go 1.25 o superior"),
    ], [1.9, 4.6], nota="Elaboración propia a partir de la documentación del proyecto (securego, s. f.) y de la ejecución de la herramienta.",
        negrita_primera_col=True, sombrear_primera_col=True)
    doc.titulo(2, "1.3 Severidad y confianza")
    doc.parrafo(
        "Cada hallazgo de gosec se califica en dos ejes independientes. La severidad (alta, media o baja) indica qué tan grave sería el "
        "problema si fuera real. La confianza (alta, media o baja) indica qué tan seguro está gosec de que el hallazgo corresponde a un "
        "problema real. Un hallazgo de severidad media y confianza baja merece una revisión humana antes de actuar, mientras que uno de "
        "severidad alta y confianza alta debe atenderse primero. Esta distinción guía la priorización de las secciones 4 y 8.")


def seccion2(doc):
    doc.titulo(1, "2. Configuración del entorno")
    doc.titulo(2, "2.1 Equipo y requisitos")
    doc.parrafo(
        "La instalación y los análisis se realizaron en un equipo con Windows 11 y Go 1.27.2, sin Docker, sin cuentas en servicios "
        "externos y sin permisos de administrador. La {T:entorno} reúne los datos del entorno exacto, registrados por el script de "
        "escaneo para que el análisis sea repetible.")
    tab(doc, "entorno", "Entorno de las ejecuciones", ["Elemento", "Valor"], [
        ("Sistema operativo", "Windows 11 Home"),
        ("Go", "1.27.2 (windows/amd64)"),
        ("gosec", "v2.30.0"),
        ("Aplicación vulnerable", "GoVWA, commit 4058f79"),
        ("Código analizado", "20 archivos y 1.569 líneas"),
        ("Fecha de las ejecuciones", "23 de octubre de 2026"),
        ("Requisitos adicionales", "Ninguno"),
    ], [2.2, 4.3], nota="Datos del archivo resultados/entorno.txt y de la ejecución de gosec.",
        negrita_primera_col=True, sombrear_primera_col=True)
    doc.titulo(2, "2.2 Instalación")
    doc.parrafo("La instalación consiste en un solo comando que descarga y compila la herramienta desde su código fuente:")
    codigo(doc, ["go install github.com/securego/gosec/v2/cmd/gosec@latest"])
    doc.parrafo(
        "La compilación tardó aproximadamente 1 minuto y 50 segundos y dejó el ejecutable, de unos 97 MB, en la carpeta de binarios de "
        "Go. No fue necesario instalar nada más. Un detalle práctico es que, al instalar desde el código fuente, el comando "
        "gosec -version responde «dev» en lugar del número de versión. La versión real se consulta con el siguiente comando, que "
        "reportó la v2.30.0:")
    codigo(doc, ["go version -m gosec.exe"])
    doc.parrafo(
        "Por eso conviene registrar la versión real en cualquier informe de análisis; también explica que el resumen de los resultados "
        "de texto muestre «Gosec: dev».")
    doc.titulo(2, "2.3 Objetivos analizados")
    doc.parrafo(
        "Se analizaron dos objetivos. El primero fue GoVWA (Go Vulnerable Web Application), una aplicación web escrita en Go con "
        "vulnerabilidades introducidas a propósito para aprender a detectarlas (0c34, s. f.); se fijó el commit 4058f79 para que el "
        "análisis sea repetible. El segundo fue código propio: un módulo pequeño con los errores más comunes de Go y su versión "
        "corregida, que permite medir el efecto de aplicar las mitigaciones (sección 5).")


def seccion3(doc):
    doc.titulo(1, "3. Metodología")
    doc.titulo(2, "3.1 Procedimiento")
    doc.parrafo(
        "En cada objetivo se ejecutó gosec en varios formatos de salida y con la opción -no-fail, que evita que el código de salida "
        "distinto de cero (el que gosec devuelve cuando hay hallazgos) interrumpa los scripts. Los comandos principales, "
        "simplificados, fueron los siguientes:")
    codigo(doc, [
        "gosec -no-fail -fmt=text  -out=resultados/govwa.txt   ./objetivos/govwa/...",
        "gosec -no-fail -fmt=json  -out=resultados/govwa.json  ./objetivos/govwa/...",
        "gosec -no-fail -fmt=sarif -out=resultados/govwa.sarif ./objetivos/govwa/...",
        "gosec -no-fail -fmt=html  -out=resultados/govwa.html  ./objetivos/govwa/...",
        "gosec -no-fail -severity=high -confidence=high -fmt=text \\",
        "      -out=resultados/govwa_solo_graves.txt ./objetivos/govwa/...",
    ])
    doc.parrafo(
        "Para el código propio se añadió la opción -track-suppressions, que registra los hallazgos suprimidos con comentarios #nosec. "
        "Los archivos JSON se procesaron con scripts en Python que resumen los hallazgos por severidad, confianza, regla y archivo.")
    doc.titulo(2, "3.2 Criterios de análisis")
    doc.parrafo(
        "Los hallazgos se clasificaron por severidad y confianza. Se revisó manualmente el código de los hallazgos de severidad alta y "
        "de los de confianza baja, y cada hallazgo destacado se contrastó con la versión segura equivalente cuando el propio proyecto la "
        "incluía. En el código propio se midió el resultado antes y después de corregir.")
    doc.titulo(2, "3.3 Reproducibilidad")
    doc.parrafo(
        "Los comandos, el commit del objetivo y las versiones de las herramientas quedaron registrados. Las rutas absolutas del equipo "
        "se eliminaron de los resultados para poder compartirlos. El Anexo A reúne los pasos para repetir el análisis.")


def seccion4(doc):
    doc.titulo(1, "4. Resultados sobre GoVWA")
    doc.titulo(2, "4.1 Panorama general")
    doc.parrafo(
        "Sobre GoVWA, gosec analizó 20 archivos con 1.569 líneas y reportó 35 hallazgos: 1 de severidad alta, 16 de severidad media y "
        "18 de severidad baja. Por confianza, 28 son de confianza alta y 7 de confianza baja. El único hallazgo de "
        "severidad alta es una inyección SQL con confianza alta, por lo que es el primero que debe atenderse. La {F:informe} muestra cómo "
        "presenta gosec sus hallazgos en el informe HTML, con filtros por severidad y confianza, y la {T:matriz} los cruza por ambos ejes.")
    figura_ancho(doc, "informe", "Informe HTML de gosec sobre GoVWA", IMG / "figura1_recorte.png",
                 "Captura real del informe generado con -fmt=html: filtros, resumen y los hallazgos G701 (inyección SQL) y G124 "
                 "(cookie sin atributos seguros).", 5.2)
    tab(doc, "matriz", "Hallazgos de gosec sobre GoVWA por severidad y confianza",
        ["Severidad", "Confianza alta", "Confianza baja", "Total"], [
            ("Alta", MATRIZ[("HIGH", "HIGH")], MATRIZ[("HIGH", "LOW")], SEV["HIGH"]),
            ("Media", MATRIZ[("MEDIUM", "HIGH")], MATRIZ[("MEDIUM", "LOW")], SEV["MEDIUM"]),
            ("Baja", MATRIZ[("LOW", "HIGH")], MATRIZ[("LOW", "LOW")], SEV["LOW"]),
            ("Total", CONF["HIGH"], CONF["LOW"], len(G)),
        ], [2.0, 1.6, 1.6, 1.3], alinear=["l", "c", "c", "c"], nota="Elaboración propia a partir de resultados/govwa.json.",
        negrita_primera_col=True, sombrear_primera_col=True)
    doc.titulo(2, "4.2 Hallazgos por regla")
    doc.parrafo(
        "Los 35 hallazgos corresponden a nueve reglas ({T:reglas} y {F:grafico}). Las de mayor impacto son la inyección SQL (G701 y G201), "
        "la criptografía débil (G401 y G501) y el HTML sin escapar (G203); las de severidad baja son las más numerosas, pero casi todas "
        "se refieren a errores sin manejar y a escritura de datos externos en los registros. La {T:mitig} resume el riesgo de cada "
        "regla y la mitigación recomendada.")
    filas = []
    for r in sorted(POR_REGLA, key=lambda r: (ORDEN_SEV[REGLA_SEV[r][0]], -POR_REGLA[r], r)):
        sev, conf, cwe = REGLA_SEV[r]
        filas.append((r, f"CWE-{cwe}", TEXTO_SEV[sev], TEXTO_SEV[conf], POR_REGLA[r], NOMBRE_REGLA[r]))
    tab(doc, "reglas", "Reglas de gosec activadas sobre GoVWA", ["Regla", "CWE", "Severidad", "Confianza", "N.°", "Hallazgo"],
        filas, [0.8, 0.95, 1.0, 1.0, 0.5, 2.25], alinear=["l", "l", "l", "l", "c", "l"],
        nota="Elaboración propia a partir de resultados/govwa.json. N.° = cantidad de hallazgos de la regla.")
    doc.figura("grafico", "Hallazgos de gosec sobre GoVWA por regla y severidad", IMG / "figura2_hallazgos_por_regla.png",
               nota="Elaboración propia a partir de resultados/govwa.json.")
    tab(doc, "mitig", "Riesgo y mitigación recomendada por regla", ["Regla", "Riesgo", "Mitigación recomendada"], [
        ("G701 y G201", "El dato del usuario se une a la instrucción SQL y cambia su significado (CWE-89).",
         "Consulta parametrizada y validación del tipo del identificador."),
        ("G203", "template.HTML desactiva el escapado: un texto con etiquetas se ejecuta en el navegador (XSS, CWE-79).",
         "Pasar texto plano a html/template; si se requiere HTML, sanitizarlo y añadir una política CSP."),
        ("G401 y G501", "MD5 es muy rápido y no usa sal: las contraseñas de una base filtrada se recuperan con facilidad.",
         "bcrypt o argon2id, y migrar los hashes existentes en el siguiente inicio de sesión."),
        ("G124", "Cookies expuestas a robo por script (sin HttpOnly), al envío por HTTP (sin Secure) o a CSRF (sin SameSite).",
         "Establecer Secure, HttpOnly y SameSite."),
        ("G112", "Un cliente que envía las cabeceras muy despacio mantiene conexiones abiertas y agota el servidor (Slowloris).",
         "Usar http.Server con ReadHeaderTimeout y los demás tiempos límite."),
        ("G706", "Un texto con saltos de línea puede falsificar entradas del registro.",
         "Registro estructurado (log/slog) y limpieza de los saltos de línea de los valores externos."),
        ("G104", "Un error ignorado deja la aplicación en un estado no previsto.",
         "Comprobar el error y responder de forma controlada."),
    ], [0.95, 3.0, 2.55], nota="Elaboración propia.", negrita_primera_col=True, sombrear_primera_col=True, tam=10)
    doc.titulo(2, "4.3 Hallazgos destacados")
    doc.titulo(3, "4.3.1 Inyección SQL (G701 y G201)")
    doc.parrafo(
        "En vulnerability/sqli/function.go, la función UnsafeQueryGetData arma la consulta con fmt.Sprintf e inserta directamente el "
        "identificador recibido del usuario (uid); después la ejecuta con DB.Query. gosec lo reporta desde dos ángulos: la regla G201 "
        "detecta el SQL armado con formato de texto (líneas 37 a 40) y la regla G701, mediante análisis de flujo de datos, detecta que el "
        "dato del usuario llega hasta Query (línea 41). La salida de texto de gosec para el segundo hallazgo fue la siguiente (se "
        "acortaron la primera línea y la sangría para ajustarlas al ancho de la página):")
    codigo(doc, [
        "[Gosec/objetivos/govwa/vulnerability/sqli/function.go:41] - G701 (CWE-89):",
        "SQL injection via taint analysis: (*database/sql.DB).Query",
        "(Confidence: HIGH, Severity: HIGH)",
        "    40:     and u.id=%s`, uid) //here is the vulnerable query",
        "  > 41:     rows, err := DB.Query(getProfileSql)",
        "    42:     if err != nil {",
    ])
    doc.parrafo(
        "El riesgo es que un atacante modifique el significado de la consulta para leer o alterar datos de otros usuarios. La mitigación "
        "es usar una consulta parametrizada, de modo que el dato viaje aparte de la instrucción SQL, y validar que el identificador sea "
        "numérico. El propio proyecto incluye junto a la función vulnerable una versión segura, SafeQueryGetData, que usa un marcador (?) "
        "y db.Prepare (líneas 66 a 75); gosec no la marcó, lo que sirve como control de que la regla no genera alarmas sobre el código "
        "corregido.")
    doc.titulo(3, "4.3.2 Contraseñas protegidas con MD5 (G401 y G501)")
    doc.parrafo(
        "En user/user.go, el inicio de sesión calcula el MD5 de la contraseña (línea 88; la función Md5Sum ocupa las líneas 159 y 160) "
        "para compararlo con lo almacenado. gosec reporta la debilidad con dos reglas: G401, que detecta el uso de MD5 (CWE-328), y G501, "
        "que detecta la importación de crypto/md5 (CWE-327). El mismo patrón aparece en idor.go y csa.go, de modo que cada regla suma "
        "tres hallazgos y, entre las dos, seis. MD5 se calcula en nanosegundos y aquí no hay sal: si la base de datos se filtra, las "
        "contraseñas comunes se recuperan casi de inmediato. La mitigación es usar bcrypt o argon2id y migrar los hashes existentes "
        "aprovechando el inicio de sesión: al validar una contraseña correcta con el hash antiguo, se guarda de nuevo con el algoritmo "
        "nuevo.")
    doc.titulo(3, "4.3.3 HTML sin escapar (G203)")
    doc.parrafo(
        "En vulnerability/xss/xss.go, el término de búsqueda se toma de r.FormValue(\"term\") (línea 39) y se convierte con template.HTML "
        "(líneas 62 y 63). Esa conversión le indica a Go que el texto ya es HTML seguro y desactiva el escapado automático de "
        "html/template. gosec reporta seis hallazgos G203 (cinco en xss.go y uno en util/template.go) con severidad media y confianza "
        "baja, porque no puede saber si el valor proviene del usuario. La revisión manual confirmó que en xss.go sí proviene del "
        "usuario, por lo que es un riesgo real de XSS (CWE-79): es el ejemplo de por qué la confianza importa, ya que lo que gosec "
        "califica con confianza baja debe verificarse antes de descartarlo. La aplicación intenta filtrar el texto con una función propia "
        "(removeScriptTag), pero filtrar con listas negras es frágil; lo correcto es pasar texto plano y dejar que html/template escape "
        "la salida.")
    doc.titulo(2, "4.4 Otros hallazgos")
    doc.parrafo(
        "Las demás reglas son de severidad media o baja. La regla G124 señala dos cookies creadas sin atributos seguros (util/cookie.go, "
        "líneas 32 y 48). La regla G112 indica que el servidor de app.go (líneas 71 a 74) no configura ReadHeaderTimeout y es vulnerable "
        "a conexiones lentas. La regla G706 reporta once escrituras de valores externos en los registros (seis en user/session/session.go, "
        "tres en util/middleware/middleware.go, una en xss.go y una en function.go). La regla G104 reporta siete errores sin manejar "
        "(cuatro en idor.go, dos en util/template.go y uno en middleware.go).")


def seccion5(doc):
    doc.titulo(1, "5. Código propio: antes y después")
    doc.titulo(2, "5.1 Diseño del ejemplo")
    doc.parrafo(
        "Para demostrar el ciclo completo de detectar, corregir y volver a escanear, se escribió un módulo pequeño en dos versiones: la "
        "versión vulnerable, de 73 líneas, que reúne los errores típicos de Go, y la versión corregida, de 126 líneas, con cada error "
        "resuelto y las mismas funciones con los mismos nombres. gosec reportó 17 hallazgos en la versión vulnerable, que corresponden a "
        "15 reglas distintas, y ninguno activo en la versión corregida ({T:antes}). La {T:correcciones} detalla la corrección aplicada a "
        "cada error.")
    tab(doc, "antes", "Resultados de gosec sobre el código propio antes y después de corregir",
        ["Medida", "Vulnerable", "Corregido"], [
            ("Líneas de código", "73", "126"),
            ("Hallazgos activos", len(V_ACTIVOS), len(C_ACTIVOS)),
            ("Severidad alta", 6, 0),
            ("Severidad media", 9, 0),
            ("Severidad baja", 2, 0),
            ("Reglas distintas", 15, 0),
            ("Hallazgos suprimidos con #nosec", 0, len(C_SUPRIM)),
        ], [3.1, 1.7, 1.7], alinear=["l", "c", "c"],
        nota="Elaboración propia a partir de resultados/ejemplo_vulnerable.json y resultados/ejemplo_corregido.json.",
        negrita_primera_col=True, sombrear_primera_col=True)
    tab(doc, "correcciones", "Errores del código propio y corrección aplicada", ["Error", "Regla", "Corrección aplicada"], [
        ("SQL armado con fmt.Sprintf", "G201, G701", "Consulta parametrizada con marcador (?)"),
        ("Contraseña con MD5", "G401, G501", "bcrypt con sal y costo por defecto"),
        ("Credencial escrita en el código", "G101", "Lectura desde la variable de entorno con os.Getenv"),
        ("TLS sin verificar el certificado", "G402", "Verificación activa y TLS 1.2 como mínimo"),
        ("Token con math/rand", "G404", "crypto/rand con 32 bytes"),
        ("Cookie sin atributos seguros", "G124", "Secure, HttpOnly y SameSite estricto"),
        ("Servidor sin tiempos límite", "G114", "http.Server con tiempos límite; solo en 127.0.0.1"),
        ("Ruta escrita por el usuario", "G304, G703", "os.OpenRoot: no permite salir del directorio raíz"),
        ("Comando con texto del usuario", "G204, G702, G705", "Lista de permitidos, sin shell ni salida devuelta"),
        ("Errores sin manejar", "G104", "Comprobación del error y respuesta controlada"),
    ], [2.0, 1.25, 3.25], nota="Elaboración propia a partir del código de ejemplo-propio/vulnerable y ejemplo-propio/corregido.", tam=10)
    doc.titulo(2, "5.2 Falso positivo y supresión justificada")
    doc.parrafo(
        "En la versión corregida, el comando ping se ejecuta solo si el destino pertenece a una lista cerrada de hosts permitidos y sin "
        "pasar por una shell. Aun así, gosec siguió reportando dos hallazgos (G204 y G702) porque no entiende esa validación: es un falso "
        "positivo. Se suprimió con un comentario que lleva además la justificación escrita (aquí partida en dos líneas por el ancho de "
        "la página):")
    codigo(doc, [
        "// #nosec G204 G702 -- host validado contra una lista cerrada",
        "//   y sin pasar por una shell",
        "if err := exec.Command(\"ping\", \"-c\", \"1\", host).Run(); err != nil {",
    ])
    doc.parrafo(
        "Con la opción -track-suppressions, gosec no oculta la supresión: la registra junto con su justificación, de modo que queda "
        "auditable. El resumen de la ejecución reportó «Nosec: 1» (un comentario) y «Issues: 0», y el archivo JSON conserva los dos "
        "hallazgos suprimidos con su justificación. La supresión debe usarse con criterio: sin justificación ni revisión, un comentario "
        "#nosec puede ocultar un problema real.")


def seccion6(doc):
    doc.titulo(1, "6. Límites de la herramienta")
    doc.parrafo(
        "Un análisis estático solo ve lo que sus reglas saben reconocer. Se documentaron los siguientes límites con evidencia obtenida de "
        "las ejecuciones.")
    doc.titulo(2, "6.1 Fallos de lógica: el caso del IDOR")
    doc.parrafo(
        "GoVWA incluye en vulnerability/idor/idor.go un módulo vulnerable a referencia directa insegura a objetos (IDOR), es decir, que "
        "permite acceder a datos de otro usuario cambiando un identificador. En ese archivo gosec reportó seis hallazgos (cuatro errores "
        "sin manejar y el uso de MD5, con sus dos reglas), pero ninguno señala el IDOR. Decidir qué usuario debería ver qué dato es un "
        "asunto de lógica de negocio que ninguna regla de patrones puede conocer.")
    doc.titulo(2, "6.2 Secretos en archivos de configuración")
    doc.parrafo(
        "La contraseña de la base de datos y la clave de sesión de GoVWA están en config/config.json. gosec solo analiza archivos .go, "
        "por lo que no las vio. Detectar secretos en archivos de configuración exige herramientas específicas de búsqueda de secretos.")
    doc.titulo(2, "6.3 Credenciales en el código: la regla G101")
    doc.parrafo(
        "La regla G101 busca credenciales escritas en el código a partir del nombre de la variable y de la aleatoriedad (entropía) del "
        "valor. Para medir su alcance se escribió un experimento con cuatro variables que contienen credenciales ({T:g101}).")
    tab(doc, "g101", "Variables con credenciales marcadas por la regla G101", ["Variable", "Valor", "Configuración por defecto",
                                                                              "Con ignore_entropy"], [
        ("claveLegible", "Cadena legible de baja entropía", "No marcada", "No marcada"),
        ("claveAleatoria", "32 caracteres aleatorios", "No marcada", "No marcada"),
        ("password", "admin123", "No marcada", "Marcada"),
        ("apiToken", "32 caracteres aleatorios", "Marcada", "Marcada"),
    ], [1.4, 2.0, 1.6, 1.5], alinear=["l", "l", "c", "c"],
        nota="Elaboración propia a partir de resultados/experimento_g101.txt. Los valores son de prueba.",
        negrita_primera_col=True, sombrear_primera_col=True)
    doc.parrafo(
        "Con la configuración por defecto solo se marcó una de las cuatro variables. La variable password con un valor trivial apareció "
        "únicamente al desactivar el filtro de entropía (ignore_entropy), y las variables cuyo nombre no contiene palabras como password, "
        "token o secret (aquí, clave…) nunca se marcaron, aunque contenían credenciales. La regla se apoya, por tanto, en palabras clave "
        "en inglés y en la entropía del valor, y no debe tomarse como garantía de que no hay credenciales en el código.")
    doc.titulo(2, "6.4 Falsos positivos y confianza baja")
    doc.parrafo(
        "Los límites anteriores son falsos negativos. También hay falsos positivos, como el de los comandos validados con lista de "
        "permitidos (sección 5.2), y hallazgos de confianza baja que resultan reales, como el HTML sin escapar (sección 4.3.3). En ambos "
        "casos la decisión final corresponde a una revisión humana: gosec acota dónde mirar, pero no reemplaza el criterio del "
        "revisor.")


def seccion7(doc):
    doc.titulo(1, "7. Experiencia de uso")
    doc.parrafo(
        "La {T:experiencia} resume la valoración del grupo según los criterios de la guía: instalación, facilidad de uso, resultados e "
        "interfaz, con la evidencia que la respalda.")
    doc.parrafo(
        "Entre los aspectos favorables se destacan la rapidez del análisis, la claridad de cada hallazgo, la diversidad de formatos de "
        "salida y la posibilidad de filtrar por severidad y confianza, que facilita la priorización. En el informe HTML, además, el "
        "filtro deshabilita las categorías que no tienen hallazgos (por ejemplo, la confianza media en la {F:informe}).")
    doc.parrafo(
        "Entre los aspectos por mejorar están la versión «dev» que reporta el binario instalado desde el código fuente, el hecho de "
        "que solo analice código Go, la dependencia de la regla G101 de los nombres de variable, la necesidad de suprimir a mano los "
        "falsos positivos con comentarios #nosec y que el informe HTML requiere conexión a internet para cargar sus estilos.")
    tab(doc, "experiencia", "Valoración de la experiencia de uso de gosec", ["Criterio", "Valoración", "Evidencia"], [
        ("Instalación", "Fácil", "Un comando (go install) y unos 2 minutos de compilación. Requiere Go; no necesitó Docker, cuenta ni "
                                 "administrador."),
        ("Facilidad de uso", "Muy fácil", "gosec ./... es suficiente. El análisis completo de GoVWA (20 archivos) tardó unos 5 segundos."),
        ("Resultados", "Claros, con falsos positivos y negativos", "Cada hallazgo trae regla, CWE, severidad, confianza, línea y código. "
                                                                  "Hubo falsos positivos (G204, G702) y falsos negativos (IDOR, G101)."),
        ("Interfaz", "Terminal e informe HTML", "El informe HTML ofrece filtros por severidad y confianza; también genera JSON, SARIF y "
                                               "JUnit para integración continua."),
        ("Integración continua", "Buena", "GitHub Action oficial; el código de salida distinto de cero cuando hay hallazgos permite "
                                          "hacer fallar un pipeline, y -no-fail lo evita cuando solo se quiere informar."),
    ], [1.3, 1.5, 3.7], nota="Elaboración propia a partir de las ejecuciones descritas en las secciones 2 a 6.",
        negrita_primera_col=True, sombrear_primera_col=True)


def seccion8(doc):
    doc.titulo(1, "8. Recomendaciones")
    doc.parrafo(
        "A partir de los resultados se proponen las siguientes prácticas para adoptar gosec en un proyecto Go. Primero, integrarlo en la "
        "integración continua con la Action oficial (securego, s. f.) y exportar los resultados en formato SARIF (OASIS, 2020), que "
        "facilita su consumo desde otras plataformas. Segundo, priorizar los hallazgos de severidad alta y confianza alta, para lo que "
        "sirve el filtro -severity=high -confidence=high (en GoVWA devolvió solo la inyección SQL), y revisar manualmente los de "
        "confianza baja.")
    doc.parrafo(
        "Tercero, tratar los comentarios #nosec como una deuda: usarlos solo con justificación escrita y ejecutar gosec con "
        "-track-suppressions para que queden auditables. Cuarto, no confiar en la regla G101 como única defensa contra credenciales en "
        "el código ni en el análisis de archivos de configuración; conviene sumar una herramienta de búsqueda de secretos. Quinto, "
        "complementar el análisis estático con DAST, SCA y revisión manual, que cubren los fallos de lógica y la configuración. Por "
        "último, registrar la versión real de gosec con go version -m y fijar el commit del objetivo para que cada análisis sea "
        "repetible.")


def conclusiones(doc):
    doc.titulo(1, "Conclusiones")
    doc.parrafo(
        "gosec se instaló con un solo comando, sin Docker ni permisos de administrador, y analizó una aplicación de 20 archivos en unos "
        "5 segundos. Reportó 35 hallazgos sobre GoVWA, entre ellos la única inyección SQL de severidad alta, y sus resultados son claros "
        "porque cada uno indica regla, CWE, severidad, confianza y código.")
    doc.parrafo(
        "El ciclo con código propio mostró el valor práctico de la herramienta: los 17 hallazgos de la versión vulnerable pasaron a 0 "
        "activos tras aplicar las mitigaciones, con dos falsos positivos suprimidos de forma justificada y auditable. A la vez, la "
        "evidencia documentó sus límites: no detectó el IDOR, no revisó los secretos de config.json y la regla G101 depende del nombre de "
        "la variable y de la entropía del valor.")
    doc.parrafo(
        "En conclusión, gosec es una herramienta recomendable como primera barrera automática en proyectos Go, especialmente dentro de la "
        "integración continua, siempre que sus resultados se interpreten con la severidad y la confianza, se revisen a mano los casos "
        "dudosos y se complemente con pruebas dinámicas, análisis de dependencias y revisión del diseño.")


def referencias(doc):
    doc.titulo(1, "Referencias", salto_antes=True)
    agregar_referencia(doc.d, [("0c34. (s. f.). ", False), ("GoVWA: Go Vulnerable Web Application", True),
                               (" [Software]. GitHub. https://github.com/0c34/govwa", False)])
    agregar_referencia(doc.d, [("MITRE. (s. f.). ", False), ("Common Weakness Enumeration (CWE)", True),
                               (". https://cwe.mitre.org/", False)])
    agregar_referencia(doc.d, [("OASIS. (2020). ", False), ("Static Analysis Results Interchange Format (SARIF) Version 2.1.0", True),
                               (" (OASIS Standard). https://docs.oasis-open.org/sarif/sarif/v2.1.0/sarif-v2.1.0.html", False)])
    agregar_referencia(doc.d, [("OWASP Foundation. (s. f.). ", False), ("Static code analysis", True),
                               (". https://owasp.org/www-community/controls/Static_Code_Analysis", False)])
    agregar_referencia(doc.d, [("securego. (s. f.). ", False), ("gosec: Go security checker", True),
                               (" [Software]. GitHub. https://github.com/securego/gosec", False)])


def anexo(doc):
    doc.titulo(1, "Anexo A. Reproducción de los resultados", salto_antes=True)
    doc.parrafo(
        "Los análisis pueden repetirse desde la carpeta Gosec del repositorio del trabajo. Se requieren Go 1.25 o superior, Git, Python y "
        "un intérprete de Bash. Los siguientes comandos instalan gosec, escanean GoVWA (el script clona el commit fijado), escanean el "
        "código propio antes y después de corregir, repiten el experimento de la regla G101 y resumen los resultados:")
    codigo(doc, [
        "go install github.com/securego/gosec/v2/cmd/gosec@latest",
        "bash scripts/escanear.sh",
        "bash scripts/escanear_ejemplo.sh",
        "bash scripts/experimento_g101.sh",
        "py scripts/analizar.py",
    ])
    doc.parrafo(
        "El script analizar.py resume los hallazgos por severidad, confianza, regla y archivo; con la opción --csv escribe además "
        "el archivo resultados/hallazgos.csv. El detalle de la estructura de la carpeta se encuentra en el archivo README.md del trabajo.")


def construir(doc):
    crear_estilo_codigo(doc.d)
    portada(doc)
    listas(doc)
    introduccion(doc)
    seccion1(doc)
    seccion2(doc)
    seccion3(doc)
    seccion4(doc)
    seccion5(doc)
    seccion6(doc)
    seccion7(doc)
    seccion8(doc)
    conclusiones(doc)
    referencias(doc)
    anexo(doc)


if __name__ == "__main__":
    previo = Documento(estricto=False)       # pasada 1: solo registra la numeración de tablas y figuras
    construir(previo)
    doc = Documento(etiquetas=previo.etiquetas, estricto=True)
    construir(doc)
    doc.guardar(SALIDA)
    print(f"docx generado: {SALIDA.name} | tablas={previo.n_tabla} | figuras={previo.n_figura}")
