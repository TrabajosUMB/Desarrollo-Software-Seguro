# -*- coding: utf-8 -*-
"""Construye Gosec_SAST_Guia7.pptx (10 minutos, 11 diapositivas).

Todas las cifras salen de resultados/*.json (ejecuciones reales de gosec); el codigo citado es textual
(sangria ajustada) y lleva su archivo y numero de linea.
Uso:  py presentacion/construir_presentacion.py
"""
import json
import pathlib
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

AQUI = pathlib.Path(__file__).resolve().parent
RES = AQUI.parent / "resultados"
IMG = AQUI / "img"
LOGO = pathlib.Path.home() / "Documents" / "Universidad" / "Plantillas" / "logo_umb.png"   # logo institucional, fuera del repositorio
SALIDA = AQUI / "Gosec_SAST_Guia7.pptx"

# ---------------------------------------------------------------- datos reales
govwa = json.loads((RES / "govwa.json").read_text(encoding="utf-8"))
vul = json.loads((RES / "ejemplo_vulnerable.json").read_text(encoding="utf-8"))
cor = json.loads((RES / "ejemplo_corregido.json").read_text(encoding="utf-8"))
ISS = govwa["Issues"]
N = len(ISS)
SEV = {s: sum(1 for i in ISS if i["severity"] == s) for s in ("HIGH", "MEDIUM", "LOW")}
CONF = {c: sum(1 for i in ISS if i["confidence"] == c) for c in ("HIGH", "LOW")}
POR_REGLA = {}
for i in ISS:
    POR_REGLA.setdefault(i["rule_id"], []).append(i)
N_VUL = len(vul["Issues"])
ACTIVOS_COR = [i for i in cor["Issues"] if not i.get("suppressions")]
SUPRIMIDOS_COR = [i for i in cor["Issues"] if i.get("suppressions")]
assert (N, SEV["HIGH"], SEV["MEDIUM"], SEV["LOW"]) == (35, 1, 16, 18), "los datos cambiaron: revisar textos"
assert N_VUL == 17 and len(ACTIVOS_COR) == 0 and len(SUPRIMIDOS_COR) == 2, "el ejemplo propio cambio: revisar textos"

# ---------------------------------------------------------------- imagenes (capturas reales del informe de gosec)
informe = Image.open(IMG / "reporte_html_inicio.png")
informe.crop((385, 45, 1380, 560)).save(IMG / "recorte_portada.png")      # dos hallazgos
informe.crop((45, 45, 365, 620)).save(IMG / "recorte_filtros.png")        # filtros y resumen

# ---------------------------------------------------------------- estilo
def C(h):
    return RGBColor.from_string(h)


BG, SURFACE, SURFACE2 = "F6F4EF", "FFFEFB", "EFECE5"
INK, INK2, MUTED, LINE = "1F2328", "3D4248", "5B5F66", "DCD8CF"
ACC, ACC_SOFT = "2B5C99", "E3EAF3"
HIGH, MED, LOW = "8C2A1C", "8A5A00", "4A6580"
SANS, SEMI, MONO = "Segoe UI", "Segoe UI Semibold", "Consolas"
SEV_COLOR = {"HIGH": HIGH, "MEDIUM": MED, "LOW": LOW}
SEV_TEXTO = {"HIGH": "ALTA", "MEDIUM": "MEDIA", "LOW": "BAJA"}
TOTAL = 11

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]


def nueva(notas):
    s = prs.slides.add_slide(VACIA)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = C(BG)
    s.notes_slide.notes_text_frame.text = notas
    return s


def caja(s, x, y, w, h, fill=SURFACE, linea=LINE, radio=0.05, grosor=1):
    f = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radio else MSO_SHAPE.RECTANGLE,
                           Inches(x), Inches(y), Inches(w), Inches(h))
    if radio:
        f.adjustments[0] = radio
    if fill:
        f.fill.solid()
        f.fill.fore_color.rgb = C(fill)
    else:
        f.fill.background()
    if linea:
        f.line.color.rgb = C(linea)
        f.line.width = Pt(grosor)
    else:
        f.line.fill.background()
    f.shadow.inherit = False
    return f


def texto(s, x, y, w, h, parrafos, tam=20, color=INK, negrita=False, fuente=SANS, align=PP_ALIGN.LEFT,
          anclaje=MSO_ANCHOR.TOP, interlineado=1.15, despues=0):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = t.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anclaje
    if isinstance(parrafos, str):
        parrafos = [parrafos]
    for k, p in enumerate(parrafos):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        par.alignment = align
        par.line_spacing = interlineado
        par.space_after = Pt(despues)
        for txt, o in ([(p, {})] if isinstance(p, str) else p):
            r = par.add_run()
            r.text = txt
            f = r.font
            f.name = o.get("fuente", fuente)
            f.size = Pt(o.get("tam", tam))
            f.bold = o.get("negrita", negrita)
            f.color.rgb = C(o.get("color", color))
    return t


def etiqueta(s, x, y, w, h, txt, fill, color="FFFFFF", tam=14, negrita=True):
    caja(s, x, y, w, h, fill=fill, linea=None, radio=0.5)
    texto(s, x, y, w, h, txt, tam=tam, color=color, negrita=negrita, align=PP_ALIGN.CENTER, anclaje=MSO_ANCHOR.MIDDLE)


def titulo(s, txt, n, kicker=None):
    caja(s, 0.7, 0.58, 0.07, 0.95 if kicker else 0.62, fill=ACC, linea=None, radio=0)
    if kicker:
        texto(s, 0.95, 0.5, 11, 0.3, kicker, tam=14, color=ACC, negrita=True)
        texto(s, 0.95, 0.82, 11.6, 0.75, txt, tam=34, negrita=True, fuente=SEMI)
    else:
        texto(s, 0.95, 0.55, 11.6, 0.75, txt, tam=34, negrita=True, fuente=SEMI)
    texto(s, 0.7, 7.02, 9, 0.3, "gosec v2.30.0  ·  Guía 7  ·  Desarrollo de Software Seguro", tam=12, color=MUTED)
    texto(s, 11.6, 7.02, 1.03, 0.3, f"{n} / {TOTAL}", tam=12, color=MUTED, align=PP_ALIGN.RIGHT)


def bloque_codigo(s, x, y, w, rotulo, lineas, marcar=(), tam=14, alto_linea=0.31):
    """lineas: lista de (numero, texto). Las lineas cuyo numero esta en `marcar` van en rojo y negrita."""
    h = 0.5 + alto_linea * len(lineas) + 0.12
    caja(s, x, y, w, h, fill=SURFACE2, linea=LINE, radio=0.03)
    texto(s, x + 0.2, y + 0.1, w - 0.4, 0.3, rotulo, tam=12, color=MUTED, negrita=True)
    t = s.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.45), Inches(w - 0.4), Inches(alto_linea * len(lineas) + 0.05))
    tf = t.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    con_numeros = any(num != "" for num, _ in lineas)   # sin numeros no se reserva la columna de la izquierda
    for k, (num, txt) in enumerate(lineas):
        par = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        par.line_spacing = 1.0
        marcada = num in marcar
        r1 = par.add_run()
        r1.text = (f"{num:>3}  " if num != "" else "     ") if con_numeros else ""
        r1.font.name, r1.font.size, r1.font.color.rgb = MONO, Pt(tam), C(MUTED)
        r2 = par.add_run()
        r2.text = txt
        r2.font.name, r2.font.size = MONO, Pt(tam)
        r2.font.bold = marcada
        r2.font.color.rgb = C(HIGH if marcada else INK)
    return h


def pie_notas(partes):
    return "\n\n".join(partes)


# ================================================================ 1. Portada
s = nueva(pie_notas([
    "TIEMPO: 0:30. Presentar el tema: análisis estático de seguridad (SAST) con gosec.",
    "Decir: gosec es un escáner de código para Go. Lo probamos sobre GoVWA, una aplicación vulnerable escrita en Go, "
    "y sobre código propio. En 10 minutos: qué es, cómo se instaló, tres hallazgos, qué pasa al corregir, "
    "lo que NO vio y nuestra experiencia de uso.",
]))
s.shapes.add_picture(str(LOGO), Inches(0.7), Inches(0.55), height=Inches(0.85))
texto(s, 0.7, 1.95, 6.3, 0.4, "GUÍA 7  ·  PRUEBAS DE SEGURIDAD AUTOMATIZADAS", tam=14, color=ACC, negrita=True)
texto(s, 0.7, 2.4, 6.3, 2.4, [[("Análisis estático de seguridad (SAST) con ", {}), ("gosec", {"color": ACC})]],
      tam=40, negrita=True, fuente=SEMI, interlineado=1.05)
texto(s, 0.7, 4.55, 6.2, 0.9, "Un escáner de código fuente para Go, probado sobre una aplicación vulnerable y sobre código propio.",
      tam=19, color=INK2, interlineado=1.2)
texto(s, 0.7, 5.8, 6.3, 1.2, ["Cobos, Dilan Steven", "Rodríguez Ángel, Santiago", "Alarcón Ospina, Víctor Alfonso"],
      tam=15, color=MUTED, interlineado=1.1)
texto(s, 0.7, 6.95, 6.3, 0.3, "Universidad Manuela Beltrán  ·  Ingeniería de Software  ·  2026", tam=12, color=MUTED)
pic = s.shapes.add_picture(str(IMG / "recorte_portada.png"), Inches(7.2), Inches(1.95), width=Inches(5.5))
pic._element.nvPicPr.cNvPr.set("descr", "Informe HTML real de gosec con dos hallazgos: inyección SQL y cookie sin banderas")
pic.line.color.rgb = C(LINE)
texto(s, 7.2, 5.0, 5.5, 0.5, "Informe HTML generado por gosec sobre GoVWA (captura real).", tam=12, color=MUTED)

# ================================================================ 2. Que es
s = nueva(pie_notas([
    "TIEMPO: 1:00. Explicar qué es.",
    "SAST es revisar el código sin ejecutarlo, como revisar los planos de una casa (analogía de la guía de clase). "
    "gosec es la herramienta SAST específica de Go: la mantiene el proyecto securego, es de código abierto (Apache 2.0).",
    "Dato para decir: según su documentación, analiza el árbol sintáctico (AST) y la representación SSA del programa. "
    "En la versión 30 trae 62 reglas, cada una con su CWE, y desde hace poco análisis de flujo de datos: sigue un dato "
    "desde la petición del usuario hasta una función peligrosa (reglas G701 a G710).",
    "La salida puede ser texto, JSON, HTML, SARIF y otros: sirve para personas y para pipelines de CI.",
]))
titulo(s, "¿Qué es gosec?", 2, "HERRAMIENTA SAST PARA GO")
texto(s, 0.95, 1.85, 5.2, 1.5,
      [[("SAST", {"negrita": True}), (": revisa el código fuente sin ejecutar la aplicación, como revisar los planos antes de construir.", {})]],
      tam=21, color=INK2, interlineado=1.25)
texto(s, 0.95, 3.45, 5.2, 2.2,
      ["Código abierto, licencia Apache 2.0.", "Mantenido por el proyecto securego.", "Requiere Go 1.25 o superior.",
       "Reglas con su CWE y filtros por severidad y confianza."],
      tam=18, color=INK2, interlineado=1.2, despues=6)
tarjetas = [
    ("62 reglas", "Credenciales, SQL, comandos, criptografía, TLS, archivos, cookies y servidores HTTP."),
    ("Flujo de datos", "Sigue la entrada del usuario hasta SQL, comandos, rutas, SSRF, XSS y logs (G701 a G710)."),
    ("9 formatos", "text, json, yaml, csv, junit-xml, html, sonarqube, golint y sarif."),
    ("Listo para CI", "GitHub Action oficial (securego/gosec) y filtros para fallar el pipeline."),
]
for k, (cab, cuerpo) in enumerate(tarjetas):
    x, y = 6.6 + (k % 2) * 3.1, 1.85 + (k // 2) * 2.45
    caja(s, x, y, 2.95, 2.3)
    texto(s, x + 0.22, y + 0.2, 2.5, 0.5, cab, tam=22, color=ACC, negrita=True, fuente=SEMI)
    texto(s, x + 0.22, y + 0.78, 2.55, 1.45, cuerpo, tam=15, color=INK2, interlineado=1.15)
caja(s, 0.95, 6.0, 5.2, 0.7, fill=ACC_SOFT, linea=None, radio=0.12)
texto(s, 1.15, 6.0, 4.9, 0.7, "Limitación clave: solo analiza código Go.", tam=17, color=ACC, negrita=True, anclaje=MSO_ANCHOR.MIDDLE)

# ================================================================ 3. Entorno
s = nueva(pie_notas([
    "TIEMPO: 1:00. Configuración y entorno.",
    "Objetivo: GoVWA, una aplicación vulnerable escrita en Go (se fijó el commit 4058f79 para que el análisis sea repetible), "
    "y código propio, que la guía también permite.",
    "Instalación: un solo comando, go install. Compiló en 1 minuto 50 segundos y deja un binario de 97 MB en la carpeta go\\bin del usuario. "
    "No hizo falta Docker, ni cuenta, ni permisos de administrador.",
    "Detalle práctico: si se instala desde el código fuente, gosec -version responde 'dev'. La versión real se lee con go version -m: es la v2.30.0.",
    "Un escaneo completo de GoVWA (20 archivos, 1.569 líneas) tarda unos 5 segundos.",
]))
titulo(s, "Configuración y entorno", 3, "INSTALACIÓN Y PRIMER ESCANEO")
pasos = [
    ("1", "Objetivo", "GoVWA: aplicación web vulnerable escrita en Go (20 archivos, 1.569 líneas) y código propio."),
    ("2", "Entorno", "Windows 11, Go 1.27.2. Sin Docker, sin cuenta y sin administrador."),
    ("3", "Instalación", "Un comando. Compila gosec desde el código fuente: 1 min 50 s."),
    ("4", "Versión real", "El binario dice «dev»; go version -m revela la v2.30.0."),
    ("5", "Escaneo", "Un análisis completo de GoVWA tarda unos 5 segundos."),
]
for k, (n, cab, cuerpo) in enumerate(pasos):
    y = 1.85 + k * 0.98
    etiqueta(s, 0.95, y + 0.05, 0.5, 0.5, n, ACC, tam=18)
    texto(s, 1.65, y, 5.2, 0.35, cab, tam=18, negrita=True)
    texto(s, 1.65, y + 0.36, 5.2, 0.6, cuerpo, tam=15, color=INK2, interlineado=1.1)
bloque_codigo(s, 7.05, 1.85, 5.65, "COMANDOS USADOS", [
    ("", "# instalar"),
    ("", "go install github.com/securego/gosec/v2/cmd/gosec@latest"),
    ("", ""),
    ("", "# comprobar la versión real"),
    ("", "go version -m gosec.exe"),
    ("", ""),
    ("", "# escanear y guardar el informe"),
    ("", "gosec -fmt=json -out=govwa.json ./..."),
    ("", "gosec -fmt=html -out=govwa.html ./..."),
], tam=12, alto_linea=0.3)
caja(s, 7.05, 5.4, 5.65, 1.15, fill=ACC_SOFT, linea=None, radio=0.08)
texto(s, 7.25, 5.4, 5.25, 1.15, [[("Objetivo fijado: ", {"negrita": True, "color": ACC}),
                                  ("GoVWA, commit 4058f79. Cada resultado se puede repetir con scripts/escanear.sh.", {})]],
      tam=15, color=INK2, anclaje=MSO_ANCHOR.MIDDLE, interlineado=1.15)

# ================================================================ 4. Resultados
s = nueva(pie_notas([
    "TIEMPO: 1:00. Resultados globales.",
    f"Sobre GoVWA gosec reportó {N} hallazgos: {SEV['HIGH']} de severidad alta, {SEV['MEDIUM']} media y {SEV['LOW']} baja.",
    f"Hay dos ejes. Severidad: qué tan grave sería si es real. Confianza: qué tan seguro está gosec de que es un problema. "
    f"De los {N}, {CONF['HIGH']} tienen confianza alta y {CONF['LOW']} baja; los de confianza baja se revisan a mano.",
    "El gráfico agrupa por regla. Lo más numeroso es poco grave (inyección en logs y errores sin manejar). "
    "Lo más grave es uno solo: la inyección SQL, que veremos primero.",
]))
titulo(s, f"{N} hallazgos en 20 archivos", 4, "RESULTADOS SOBRE GOVWA")
ORDEN = [("G701", "Inyección SQL"), ("G203", "HTML sin escapar"), ("G401", "Criptografía débil"), ("G501", "Importa crypto/md5"),
         ("G124", "Cookie sin banderas"), ("G201", "SQL con formato"), ("G112", "Sin ReadHeaderTimeout"),
         ("G706", "Inyección en logs"), ("G104", "Error sin manejar")]
assert {r for r, _ in ORDEN} == set(POR_REGLA), "hay reglas nuevas en los resultados"
datos = CategoryChartData()
datos.categories = [f"{r}  {n}  ·  {SEV_TEXTO[POR_REGLA[r][0]['severity']]}" for r, n in ORDEN]
datos.add_series("Hallazgos", [len(POR_REGLA[r]) for r, _ in ORDEN])
graf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(1.75), Inches(8.1), Inches(5.1), datos).chart
graf.has_legend = False
graf.has_title = False
graf.font.name, graf.font.size, graf.font.color.rgb = SANS, Pt(14), C(INK2)
graf.category_axis.reverse_order = True
graf.category_axis.format.line.color.rgb = C(LINE)
graf.category_axis.tick_labels.font.size = Pt(14)
graf.value_axis.visible = False
graf.value_axis.has_major_gridlines = False
graf.value_axis.maximum_scale = 13
graf.plots[0].gap_width = 45
graf.plots[0].has_data_labels = True
etq = graf.plots[0].data_labels
etq.font.size, etq.font.bold, etq.font.name = Pt(15), True, SANS
etq.font.color.rgb = C(INK)
etq.position = XL_LABEL_POSITION.OUTSIDE_END
for k, (r, _) in enumerate(ORDEN):
    p = graf.series[0].points[k]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = C(SEV_COLOR[POR_REGLA[r][0]["severity"]])
for k, (cab, n, col, sub) in enumerate([
        ("ALTA", SEV["HIGH"], HIGH, "inyección SQL"),
        ("MEDIA", SEV["MEDIUM"], MED, "HTML, cripto, cookies, SQL"),
        ("BAJA", SEV["LOW"], LOW, "logs y errores sin manejar")]):
    y = 1.85 + k * 1.3
    caja(s, 9.2, y, 3.45, 1.15)
    etiqueta(s, 9.4, y + 0.28, 0.95, 0.6, str(n), col, tam=26)
    texto(s, 10.5, y + 0.2, 2.1, 0.4, f"Severidad {cab}", tam=16, negrita=True)
    texto(s, 10.5, y + 0.58, 2.1, 0.5, sub, tam=13, color=MUTED, interlineado=1.05)
caja(s, 9.2, 5.75, 3.45, 1.1, fill=ACC_SOFT, linea=None, radio=0.08)
texto(s, 9.4, 5.75, 3.05, 1.1, [[("Confianza: ", {"negrita": True, "color": ACC}),
                                  (f"{CONF['HIGH']} alta y {CONF['LOW']} baja. La baja se revisa a mano.", {})]],
      tam=14, color=INK2, anclaje=MSO_ANCHOR.MIDDLE, interlineado=1.1)

# ================================================================ 5-7. Hallazgos
def hallazgo(n, kicker, titulo_txt, sev, conf, regla, cwe, archivo, codigo, marcar, corr_rotulo, corr_codigo,
             riesgo, mitigacion, nota_inferior, notas):
    s = nueva(notas)
    titulo(s, titulo_txt, n, kicker)
    h1 = bloque_codigo(s, 0.7, 1.8, 7.5, f"GOSEC MARCÓ  ·  {archivo}", codigo, marcar)
    bloque_codigo(s, 0.7, 1.8 + h1 + 0.2, 7.5, corr_rotulo, corr_codigo)
    x = 8.5
    etiqueta(s, x, 1.8, 2.05, 0.46, f"SEVERIDAD {SEV_TEXTO[sev]}", SEV_COLOR[sev], tam=13)
    etiqueta(s, x + 2.15, 1.8, 2.0, 0.46, f"CONFIANZA {SEV_TEXTO[conf]}", SEV_COLOR[conf], tam=13)
    texto(s, x, 2.4, 4.15, 0.35, f"{regla}  ·  {cwe}", tam=15, color=MUTED, negrita=True)
    texto(s, x, 2.9, 4.15, 0.35, "Riesgo", tam=19, color=ACC, negrita=True, fuente=SEMI)
    texto(s, x, 3.28, 4.15, 1.7, riesgo, tam=15, color=INK2, interlineado=1.15, despues=4)
    texto(s, x, 5.0, 4.15, 0.35, "Mitigación", tam=19, color=ACC, negrita=True, fuente=SEMI)
    texto(s, x, 5.38, 4.15, 1.5, mitigacion, tam=15, color=INK2, interlineado=1.15, despues=4)
    if nota_inferior:
        caja(s, 0.7, 6.28, 7.5, 0.6, fill=ACC_SOFT, linea=None, radio=0.1)
        texto(s, 0.9, 6.28, 7.1, 0.6, nota_inferior, tam=14, color=ACC, negrita=True, anclaje=MSO_ANCHOR.MIDDLE)
    return s


hallazgo(5, "HALLAZGO 1  ·  EL ÚNICO DE SEVERIDAD ALTA", "Inyección SQL", "HIGH", "HIGH",
         "G701 y G201", "CWE-89", "vulnerability/sqli/function.go",
         [(37, "getProfileSql := fmt.Sprintf(`SELECT p.user_id, p.full_name,"),
          (38, "    p.city, p.phone_number FROM Profile as p, Users as u"),
          (39, "    where p.user_id = u.id"),
          (40, "    and u.id=%s`, uid) //here is the vulnerable query"),
          (41, "rows, err := DB.Query(getProfileSql)")],
         (40, 41),
         "CORRECCIÓN DEL PROPIO REPOSITORIO  ·  SafeQueryGetData",
         [(66, "and u.id=?`"),
          (69, "stmt, err := DB.Prepare(getProfileSql) //prepare statement"),
          (75, "err = stmt.QueryRow(uid).Scan(&p.Uid, &p.Name, ...)")],
         ["El identificador del usuario se pega dentro de la instrucción SQL.",
          "Un atacante cambia su significado y lee o altera datos de otros usuarios."],
         ["Consulta parametrizada: el dato viaja aparte (?).", "Validar el tipo: el id debe ser numérico."],
         "Control: la versión segura, en el mismo archivo, no fue marcada.",
         pie_notas([
             "TIEMPO: 1:15. Hallazgo 1: inyección SQL, el único de severidad ALTA con confianza ALTA.",
             "Mostrar el código: la consulta se arma con fmt.Sprintf y el id que llega del usuario; luego se ejecuta con DB.Query. "
             "gosec lo reporta dos veces desde dos ángulos: G201 (SQL armado con formato) y G701 (análisis de flujo de datos: "
             "el dato del usuario llega hasta Query).",
             "Riesgo: el atacante cambia el sentido de la consulta y lee o modifica datos de otros usuarios.",
             "Mitigación: consulta parametrizada. Dato interesante: el propio GoVWA trae al lado una función segura, SafeQueryGetData, "
             "con DB.Prepare y un signo de interrogación. gosec NO la marcó: sirve como control de que no inventa hallazgos aquí.",
         ]))

hallazgo(6, "HALLAZGO 2  ·  SEVERIDAD MEDIA, CONFIANZA ALTA", "Contraseñas protegidas con MD5", "MEDIUM", "HIGH",
         "G401 y G501", "CWE-328 y CWE-327", "user/user.go",
         [(88, 'pass := Md5Sum(r.FormValue("password"))'),
          (159, "func Md5Sum(text string) string {"),
          (160, "    hasher := md5.New()")],
         (88, 160),
         "NUESTRA CORRECCIÓN  ·  ejemplo-propio/corregido",
         [("", "h, err := bcrypt.GenerateFromPassword("),
          ("", "    []byte(clave), bcrypt.DefaultCost)"),
          ("", "// al iniciar sesión:"),
          ("", "// bcrypt.CompareHashAndPassword(h, ...)")],
         ["MD5 es muy rápido y aquí no lleva sal.", "Con una base de datos filtrada, las contraseñas comunes se recuperan en segundos."],
         ["bcrypt o argon2id: lentos a propósito y con sal.", "Migrar los hashes viejos al iniciar sesión."],
         "La misma debilidad en 3 archivos: 6 hallazgos (user.go, idor.go y csa.go).",
         pie_notas([
             "TIEMPO: 1:00. Hallazgo 2: contraseñas con MD5.",
             "En user.go, línea 88, el inicio de sesión calcula MD5 de la contraseña y la compara con lo guardado. "
             "gosec lo reporta con dos reglas: G401, que detecta el uso de MD5 en el código, y G501, que detecta el import de crypto/md5. "
             "Es la misma debilidad repetida en tres archivos (user.go, idor.go y csa.go), por eso salen 6 hallazgos.",
             "Riesgo: MD5 se calcula en nanosegundos y aquí no hay sal; si se filtra la base, las contraseñas comunes se recuperan casi al instante.",
             "Mitigación: bcrypt o argon2id. Esa es la corrección que aplicamos en nuestro código propio.",
         ]))

hallazgo(7, "HALLAZGO 3  ·  CONFIANZA BAJA: REQUIERE REVISIÓN HUMANA", "HTML sin escapar (XSS)", "MEDIUM", "LOW",
         "G203", "CWE-79", "vulnerability/xss/xss.go",
         [(39, 'term := r.FormValue("term")'),
          (61, 'vuln := fmt.Sprintf("<b>%s</b>", term)'),
          (62, 'data["value"] = template.HTML(value)'),
          (63, 'data["term"] = template.HTML(vuln)')],
         (62, 63),
         "NUESTRA CORRECCIÓN",
         [("", 'data["term"] = term   // html/template lo escapa solo'),
          ("", "// HTML propio: solo constantes, o sanitizado")],
         ["template.HTML le dice a Go que el texto ya es seguro y salta el escapado automático.",
          "Un término con una etiqueta de script se ejecutaría en el navegador de la víctima."],
         ["Pasar texto plano: html/template escapa solo.", "Si hace falta HTML, sanitizarlo y añadir una política CSP."],
         "6 hallazgos G203 en 2 archivos. gosec duda (confianza baja): aquí sí son reales.",
         pie_notas([
             "TIEMPO: 1:00. Hallazgo 3: HTML sin escapar, riesgo de XSS.",
             "En xss.go, el término sale de r.FormValue, o sea lo escribe el usuario, y se convierte con template.HTML. "
             "Esa conversión le dice a Go que el texto ya es HTML seguro y apaga el escapado automático de html/template.",
             "gosec lo marca con severidad media pero confianza BAJA: no puede saber si el dato viene del usuario. "
             "Aquí lo revisamos a mano y sí es real. Es el ejemplo de por qué la confianza importa: lo de confianza baja se verifica antes de actuar.",
             "Nota: la app intenta filtrar con una función removeScriptTag; filtrar con listas negras es frágil, lo correcto es escapar la salida.",
         ]))

# ================================================================ 8. Antes y despues
def tabla(s, x, y, anchos, alto_fila, encabezado, filas, tam=13):
    forma = s.shapes.add_table(len(filas) + 1, len(anchos), Inches(x), Inches(y), Inches(sum(anchos)), Inches(alto_fila * (len(filas) + 1)))
    t = forma.table
    for k, a in enumerate(anchos):
        t.columns[k].width = Inches(a)
    for r in range(len(filas) + 1):
        t.rows[r].height = Inches(alto_fila)
        for c in range(len(anchos)):
            celda = t.cell(r, c)
            celda.margin_left = celda.margin_right = Inches(0.1)
            celda.margin_top = celda.margin_bottom = Inches(0.03)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            celda.fill.solid()
            if r == 0:
                celda.fill.fore_color.rgb = C(ACC)
                valor = encabezado[c]
            else:
                celda.fill.fore_color.rgb = C(SURFACE if r % 2 else SURFACE2)
                valor = filas[r - 1][c]
            tf = celda.text_frame
            tf.word_wrap = True
            par = tf.paragraphs[0]
            par.line_spacing = 1.0
            run = par.add_run()
            run.text = valor
            run.font.name = SEMI if (r == 0 or c == 0) else SANS
            run.font.size = Pt(tam)
            run.font.bold = r == 0
            run.font.color.rgb = C("FFFFFF" if r == 0 else INK)
    return t


s = nueva(pie_notas([
    "TIEMPO: 1:15. Antes y después con código propio.",
    f"La guía permite analizar código propio. Escribimos un módulo pequeño con los errores típicos de Go ({N_VUL} hallazgos, "
    "15 reglas distintas) y su versión corregida. Después de corregir, gosec reporta 0 hallazgos activos.",
    "La tabla resume nueve correcciones: consulta parametrizada, bcrypt, credencial desde variable de entorno, TLS con verificación, "
    "crypto/rand, cookie con banderas, servidor con tiempos límite, os.OpenRoot para rutas, y para el comando una lista de permitidos sin shell.",
    "Punto honesto: en el comando, gosec sigue avisando (G204 y G702) aunque ya validamos con lista de permitidos, porque no entiende esa validación. "
    "Se suprime con un comentario #nosec que lleva la justificación escrita, y gosec lo registra: 2 suprimidos, 0 activos.",
]))
titulo(s, "Antes y después: código propio", 8, "EL CICLO COMPLETO: DETECTAR, CORREGIR, VOLVER A ESCANEAR")
caja(s, 0.7, 1.85, 3.5, 4.95)
texto(s, 0.7, 2.0, 3.5, 1.0, str(N_VUL), tam=64, color=HIGH, negrita=True, fuente=SEMI, align=PP_ALIGN.CENTER)
texto(s, 0.9, 3.05, 3.1, 0.9, "hallazgos en el código vulnerable (6 altos, 9 medios y 2 bajos)", tam=14, color=INK2, align=PP_ALIGN.CENTER, interlineado=1.1)
texto(s, 0.7, 3.95, 3.5, 0.6, "↓", tam=30, color=MUTED, align=PP_ALIGN.CENTER)
texto(s, 0.7, 4.4, 3.5, 1.0, "0", tam=64, color=ACC, negrita=True, fuente=SEMI, align=PP_ALIGN.CENTER)
texto(s, 0.9, 5.45, 3.1, 1.2, ["hallazgos activos tras corregir", "(2 suprimidos con #nosec y justificación escrita)"],
      tam=14, color=INK2, align=PP_ALIGN.CENTER, interlineado=1.1, despues=4)
tabla(s, 4.5, 1.85, [3.65, 4.5], 0.5, ["Error (regla de gosec)", "Corrección aplicada"], [
    ("SQL armado a mano (G201, G701)", "Consulta parametrizada: db.Query(sql, id)"),
    ("Contraseñas con MD5 (G401, G501)", "bcrypt con sal y costo"),
    ("Credencial en el código (G101)", "Se lee con os.Getenv, nunca del fuente"),
    ("TLS sin verificar (G402)", "Verificación activa y TLS 1.2 como mínimo"),
    ("Tokens con math/rand (G404)", "crypto/rand"),
    ("Cookie sin banderas (G124)", "Secure, HttpOnly y SameSite"),
    ("Servidor sin límites (G114)", "http.Server con tiempos límite, solo en 127.0.0.1"),
    ("Ruta del usuario (G304, G703)", "os.OpenRoot: no puede salir de la carpeta"),
    ("Comando del usuario (G204, G702, G705)", "Lista de permitidos, sin shell ni eco de salida"),
], tam=13)

# ================================================================ 9. Lo que no vio
s = nueva(pie_notas([
    "TIEMPO: 1:15. Lo que gosec no vio. Esta es la parte más útil para decidir cómo usarlo.",
    "1) IDOR: es un fallo de lógica (acceder a datos de otro usuario cambiando un identificador). En idor.go gosec marcó el MD5 y errores sin manejar, "
    "pero no el IDOR en sí: ninguna regla de patrón puede saber qué usuario debería ver qué.",
    "2) Secretos en config.json: la contraseña de la base de datos y la clave de sesión están en un JSON; gosec solo lee archivos .go.",
    "3) Hicimos un experimento con la regla G101 (credenciales escritas en el código): de cuatro variables con credenciales, "
    "por defecto marcó una. La variable password con valor admin123 solo aparece si se configura ignore_entropy, y las llamadas clave... "
    "nunca, porque la regla busca nombres en inglés como password, token o secret. Resultados reales en resultados/experimento_g101.txt.",
    "4) Falsos positivos: tras validar con una lista de permitidos, G204 y G702 siguen avisando. Se suprimen con #nosec y una justificación.",
    "Conclusión: SAST se complementa con DAST, SCA y revisión manual, como dice la guía de clase.",
]))
titulo(s, "Lo que gosec no vio", 9, "LÍMITES, CON EVIDENCIA")
limites = [
    ("IDOR", "Es un fallo de lógica. En idor.go gosec marcó el MD5 y errores sin manejar, pero no el acceso a datos de otro usuario."),
    ("Secretos en config.json", "La contraseña de la base y la clave de sesión están en un JSON. gosec solo analiza archivos .go."),
    ("Credenciales: depende del nombre", "Con 4 variables con credenciales, G101 marcó 1. password = \"admin123\" exige ignore_entropy y las variables llamadas clave... nunca."),
    ("Falsos positivos", "Tras validar con lista de permitidos, G204 y G702 siguen avisando. Se suprimen con #nosec y justificación; gosec lo registra."),
]
for k, (cab, cuerpo) in enumerate(limites):
    x, y = 0.7 + (k % 2) * 6.1, 1.85 + (k // 2) * 2.25
    caja(s, x, y, 5.95, 2.1)
    texto(s, x + 0.25, y + 0.2, 5.5, 0.4, cab, tam=20, color=ACC, negrita=True, fuente=SEMI)
    texto(s, x + 0.25, y + 0.72, 5.45, 1.3, cuerpo, tam=15, color=INK2, interlineado=1.2)
caja(s, 0.7, 6.4, 12.05, 0.5, fill=ACC_SOFT, linea=None, radio=0.1)
texto(s, 0.9, 6.4, 11.65, 0.5, "Por eso SAST se complementa con DAST, SCA y revisión manual.", tam=16, color=ACC, negrita=True, anclaje=MSO_ANCHOR.MIDDLE)

# ================================================================ 10. Experiencia y conclusiones
s = nueva(pie_notas([
    "TIEMPO: 1:00. Experiencia de uso y conclusiones.",
    "Instalación: fácil. Un comando; tarda unos 2 minutos porque compila, y requiere Go. No necesitó Docker, a diferencia de la herramienta del ejemplo.",
    "Uso: muy fácil. gosec ./... y listo; unos 5 segundos para GoVWA.",
    "Resultados: claros y accionables, cada hallazgo trae regla, CWE, severidad, confianza, línea y el código. Con falsos positivos y falsos negativos, como cualquier SAST.",
    "Interfaz: es de terminal, pero genera un informe HTML con filtros por severidad y confianza (la captura de la derecha es real), y JSON, SARIF y JUnit para CI.",
    "Cierre: en 5 segundos encontró un SQLi real y 34 hallazgos más; corregir el código los llevó a cero. "
    "Recomendación: usarlo en CI con la Action oficial, priorizar severidad alta con confianza alta, y complementarlo con DAST y revisión.",
]))
titulo(s, "Experiencia de uso y conclusiones", 10, "VALORACIÓN DEL GRUPO")
tabla(s, 0.7, 1.85, [1.9, 5.9], 0.62, ["Criterio", "Nuestra experiencia"], [
    ("Instalación", "Fácil: un comando (go install), unos 2 min compilando. Requiere Go; no necesita Docker."),
    ("Facilidad de uso", "Muy fácil: gosec ./... y listo. Un análisis completo de GoVWA tarda unos 5 s."),
    ("Resultados", "Claros: regla, CWE, severidad, confianza, línea y código. Con falsos positivos y negativos."),
    ("Interfaz", "Terminal e informe HTML con filtros; JSON, SARIF y JUnit para CI."),
], tam=14)
texto(s, 0.7, 5.3, 7.8, 0.35, "Conclusiones", tam=19, color=ACC, negrita=True, fuente=SEMI)
texto(s, 0.7, 5.72, 7.8, 1.1,
      ["En 5 s encontró un SQLi real y 34 hallazgos más; corregir el código los llevó a 0.",
       "Úsalo en CI (Action oficial) y atiende primero severidad alta con confianza alta.",
       "No reemplaza DAST ni la revisión: no vio el IDOR ni los secretos de config.json."],
      tam=14, color=INK2, interlineado=1.1, despues=3)
pic = s.shapes.add_picture(str(IMG / "recorte_filtros.png"), Inches(9.5), Inches(1.85), width=Inches(2.6))
pic._element.nvPicPr.cNvPr.set("descr", "Panel de filtros y resumen del informe HTML de gosec")
pic.line.color.rgb = C(LINE)
texto(s, 8.9, 6.6, 3.8, 0.3, "Informe HTML de gosec (captura real)", tam=12, color=MUTED, align=PP_ALIGN.CENTER)

# ================================================================ 11. Preguntas
s = nueva("TIEMPO: resto. Preguntas. Todo es reproducible desde la carpeta Guia 7/Gosec del repositorio: scripts, resultados y código propio.")
texto(s, 0.7, 2.5, 11.9, 1.4, "Preguntas y dudas", tam=54, negrita=True, fuente=SEMI, align=PP_ALIGN.CENTER)
caja(s, 3.4, 4.1, 6.5, 0.07, fill=ACC, linea=None, radio=0)
texto(s, 1.7, 4.5, 9.9, 1.4,
      ["Todo es reproducible: carpeta Guia 7/Gosec del repositorio.",
       "Scripts, resultados reales, código propio y esta presentación."],
      tam=19, color=INK2, align=PP_ALIGN.CENTER, interlineado=1.25, despues=4)
texto(s, 0.7, 7.02, 9, 0.3, "gosec v2.30.0  ·  Guía 7  ·  Desarrollo de Software Seguro", tam=12, color=MUTED)
texto(s, 11.6, 7.02, 1.03, 0.3, f"{TOTAL} / {TOTAL}", tam=12, color=MUTED, align=PP_ALIGN.RIGHT)

assert len(prs.slides) == TOTAL, f"se esperaban {TOTAL} diapositivas y hay {len(prs.slides)}"
prs.core_properties.author = "Cobos, Dilan Steven; Rodríguez Ángel, Santiago; Alarcón Ospina, Víctor Alfonso"
prs.core_properties.last_modified_by = prs.core_properties.author
prs.core_properties.title = "Análisis estático de seguridad (SAST) con gosec"
prs.save(SALIDA)
print("guardado:", SALIDA, "| diapositivas:", len(prs.slides))
